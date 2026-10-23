"""Reference-only input/output. Unselected entries stay verbatim, in input order."""
from io import BytesIO
from pathlib import Path
import zipfile
from docx import Document


def read_docx(data):
    if len(data) > 5_000_000: raise ValueError('DOCX upload exceeds 5 MB.')
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 25_000_000:
                raise ValueError('DOCX expanded content exceeds 25 MB.')
        document = Document(BytesIO(data))
    except (zipfile.BadZipFile, KeyError):
        raise ValueError('Invalid DOCX file.') from None
    return [p.text.strip() for p in document.paragraphs if p.text.strip()]


def read_input(path):
    p = Path(path)
    if p.suffix.lower() == '.docx': return read_docx(p.read_bytes())
    return [line.strip() for line in p.read_text().splitlines() if line.strip()]


def format_record(record):
    # Simple transparent reference format; not a named citation style.
    parts = ['; '.join(record['authors']), f"({record['year']})" if record['year'] else '', record['title'], record['venue']]
    text = '. '.join(p.strip(' .') for p in parts if p) + '.'
    if record['doi']: text += ' https://doi.org/' + record['doi']
    elif record['pmid']: text += ' https://pubmed.ncbi.nlm.nih.gov/' + record['pmid'] + '/'
    return text


def selected_lines(report, selections):
    valid_keys = {str(row['index']) for row in report['results']}
    if set(selections) - valid_keys: raise ValueError('Unknown reference selection.')
    lines = []
    for row in report['results']:
        choice = selections.get(str(row['index']))
        if choice is None:
            lines.append(row['input']['raw'])
            continue
        candidate = next((c for c in row['candidates'] if c['id'] == str(choice)), None)
        if candidate is None: raise ValueError('Unknown candidate selection.')
        lines.append(format_record(candidate['record']))
    return lines


def export_document(report, selections, kind):
    lines = selected_lines(report, selections)
    if kind == 'txt': return ('\n\n'.join(lines) + '\n').encode()
    if kind != 'docx': raise ValueError('Export format must be txt or docx.')
    document = Document()
    for line in lines: document.add_paragraph(line)
    output = BytesIO(); document.save(output)
    return output.getvalue()
